import os
import sys
import json
import csv
import re
import argparse
import requests
from bs4 import BeautifulSoup
import openpyxl
from datetime import datetime
from utils.email_sender import send_email

# Constants
BASE_URL = "https://efpublic.elections.ab.ca"
LIST_URL = f"{BASE_URL}/efRTPAs.cfm?MID=TPAS_TR&TPATYPE=R"
XLS_URL_TEMPLATE = f"{BASE_URL}/efOFSRTPAWeeklyXLS.cfm?EVENTID={{event_id}}&ACCOUNTID={{account_id}}&TPAID={{tpa_id}}"
STATE_FILE = "data/state.json"
CSV_FILE = "data/referendum_donations.csv"
NOTABLE_THRESHOLD = 1000.0  # Configure individual donation threshold for alerts

def load_state():
    if os.path.exists(STATE_FILE):
        try:
            with open(STATE_FILE, 'r', encoding='utf-8') as f:
                return json.load(f)
        except Exception as e:
            print(f"Warning: Could not read state file: {e}. Starting fresh.")
    return {"last_processed_date": "", "groups": {}}

def save_state(state):
    os.makedirs(os.path.dirname(STATE_FILE), exist_ok=True)
    with open(STATE_FILE, 'w', encoding='utf-8') as f:
        json.dump(state, f, indent=2)
    print(f"State saved to {STATE_FILE}")

def load_previous_donations():
    """
    Loads previous individual donations from the CSV database.
    Returns: dict mapping group_id -> { donation_key_tuple: [list of report_dates] }
    """
    old_donations = {}
    if not os.path.exists(CSV_FILE):
        return old_donations
        
    try:
        with open(CSV_FILE, mode='r', encoding='utf-8') as f:
            reader = csv.DictReader(f)
            for row in reader:
                group_id = row['group_id']
                contributor = row['contributor']
                location = row['location']
                try:
                    cash = float(row['cash']) if row['cash'] else 0.0
                    valued = float(row['valued']) if row['valued'] else 0.0
                    total = float(row['total']) if row['total'] else 0.0
                except ValueError:
                    cash, valued, total = 0.0, 0.0, 0.0
                    
                report_date = row['report_date']
                
                # Skip the small donations summaries, they are managed cumulatively
                if contributor == "[Summary - Contributions $250 and under]":
                    continue
                    
                key = (contributor, location, cash, valued, total)
                
                if group_id not in old_donations:
                    old_donations[group_id] = {}
                if key not in old_donations[group_id]:
                    old_donations[group_id][key] = []
                old_donations[group_id][key].append(report_date)
    except Exception as e:
        print(f"Warning: Error reading previous CSV donations: {e}")
        
    return old_donations

def scrape_main_page():
    """
    Scrapes the main list page to discover groups and get the latest update date.
    Returns: (update_date, groups_dict)
    """
    print(f"Fetching main page: {LIST_URL}")
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Referendum Tracker'}
    response = requests.get(LIST_URL, headers=headers, timeout=30)
    response.raise_for_status()
    
    soup = BeautifulSoup(response.text, 'lxml')
    
    # 1. Parse update date from header link
    update_date = None
    header_links = soup.find_all('a', class_='CHLink')
    for link in header_links:
        text = link.get_text()
        if "Contributions to" in text:
            match = re.search(r"Contributions to\s+([A-Za-z]{3}\s+\d{1,2},\s+\d{4})", text)
            if match:
                update_date = match.group(1).strip()
                break
                
    if not update_date:
        # Fallback to date in intro paragraph if header parsing fails
        intro_text = soup.get_text()
        match = re.search(r"Contributions to\s+([A-Za-z]{3}\s+\d{1,2},\s+\d{4})", intro_text)
        if match:
            update_date = match.group(1).strip()
        else:
            # Fallback to current date
            update_date = datetime.now().strftime("%b %d, %Y")
            print(f"Warning: Could not parse update date from portal. Using fallback: {update_date}")
            
    # 2. Parse registered groups
    groups = {}
    group_links = soup.find_all('a', href=re.compile(r"efRTPA\.cfm\?TPAID=\d+"))
    print(f"Discovered {len(group_links)} groups on main page.")
    
    for link in group_links:
        name = link.get_text().strip()
        href = link['href']
        tpa_id = re.search(r"TPAID=(\d+)", href).group(1)
        
        td = link.find_parent('td')
        td_text = td.get_text()
        
        # Registration Date
        reg_match = re.search(r"Registered\s+([A-Za-z]{3}\s+\d{1,2},\s+\d{4})", td_text)
        reg_date = reg_match.group(1).strip() if reg_match else "Unknown"
        
        # Next TD contains report link with total contributions
        next_td = td.find_next_sibling('td')
        report_link = next_td.find('a', class_='ReportLink') if next_td else None
        
        event_id = None
        account_id = None
        total_contributions = 0.0
        
        if report_link:
            onclick = report_link.get('onclick', '')
            arg_match = re.search(r"fnShowOFSTPAContributionsWeekly\((\d+),\s*(\d+),\s*(\d+)\)", onclick)
            if arg_match:
                event_id = int(arg_match.group(1))
                account_id = int(arg_match.group(2))
                
            amount_text = report_link.get_text().strip().replace('$', '').replace(',', '')
            try:
                total_contributions = float(amount_text)
            except ValueError:
                total_contributions = 0.0
                
        groups[tpa_id] = {
            'tpa_id': tpa_id,
            'name': name,
            'registered_date': reg_date,
            'event_id': event_id,
            'account_id': account_id,
            'total_contributions': total_contributions
        }
        
    return update_date, groups

def download_and_parse_xls(event_id, account_id, tpa_id):
    """
    Downloads the weekly report Excel sheet for a group and parses contributions.
    Returns: (list of donations, small_contributions_total)
    """
    url = XLS_URL_TEMPLATE.format(event_id=event_id, account_id=account_id, tpa_id=tpa_id)
    print(f"Downloading report for TPAID={tpa_id} from {url}")
    headers = {'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) Referendum Tracker'}
    
    response = requests.get(url, headers=headers, timeout=30)
    response.raise_for_status()
    
    # Save temporarily to parse it
    temp_path = f"data/temp_{tpa_id}.xlsx"
    os.makedirs("data", exist_ok=True)
    with open(temp_path, "wb") as f:
        f.write(response.content)
        
    donations = []
    small_contributions_total = 0.0
    
    try:
        # openpyxl might throw warnings for missing default styles, which we can ignore
        wb = openpyxl.load_workbook(temp_path, data_only=True)
        sheet = wb.active
        
        # Locate header row
        header_row_idx = None
        for r in range(1, 20):
            row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 6)]
            if row_vals and row_vals[0] == 'Contributor' and row_vals[1] == 'Location':
                header_row_idx = r
                break
                
        if not header_row_idx:
            print(f"Warning: Could not find header row in XLSX for TPAID={tpa_id}. Assuming empty.")
            return [], 0.0
            
        # Parse contributor records
        for r in range(header_row_idx + 1, sheet.max_row + 1):
            row_vals = [sheet.cell(row=r, column=c).value for c in range(1, 6)]
            first_cell = row_vals[0]
            
            if first_cell is None:
                if all(cell is None for cell in row_vals):
                    continue
                continue
                
            first_cell_str = str(first_cell).strip()
            
            # Small donations summary row
            if first_cell_str.startswith("Contributions $250.00 and under:"):
                val = row_vals[4]
                try:
                    small_contributions_total = float(val) if val is not None else 0.0
                except ValueError:
                    small_contributions_total = 0.0
                continue
                
            # Footer row marking end of data
            if first_cell_str.startswith(("Contributions over $250.00:", "Grand Total:", "Total Contributions:")):
                break
                
            contributor = first_cell_str
            location = str(row_vals[1]).strip() if row_vals[1] is not None else ""
            
            try:
                cash = float(row_vals[2]) if row_vals[2] is not None else 0.0
                valued = float(row_vals[3]) if row_vals[3] is not None else 0.0
                total = float(row_vals[4]) if row_vals[4] is not None else 0.0
            except ValueError:
                cash, valued, total = 0.0, 0.0, 0.0
                
            donations.append({
                'contributor': contributor,
                'location': location,
                'cash': cash,
                'valued': valued,
                'total': total
            })
            
    finally:
        # Ensure temporary file is cleaned up
        if os.path.exists(temp_path):
            try:
                os.remove(temp_path)
            except Exception as e:
                print(f"Error removing temp file {temp_path}: {e}")
                
    return donations, small_contributions_total

def generate_email_report(update_date, groups, new_groups, notable_donations, group_changes, all_new_donations):
    """
    Generates beautiful HTML and Plain Text versions of the alert report.
    """
    current_time_str = datetime.now().strftime("%B %d, %Y at %I:%M %p")
    
    # ---------------------- PLAIN TEXT BODY ----------------------
    text_lines = [
        "ALBERTA REFERENDUM DONATIONS UPDATE",
        f"Portal Update Date: {update_date}",
        f"Report Generated: {current_time_str}",
        "=" * 60,
        ""
    ]
    
    # Overview
    text_lines.append("OVERVIEW SUMMARY")
    text_lines.append("-" * 30)
    grand_total = sum(g['total_contributions'] for g in groups.values())
    grand_change = sum(change['net_change'] for change in group_changes.values())
    text_lines.append(f"Total Advertisers: {len(groups)}")
    text_lines.append(f"Grand Total Contributions: ${grand_total:,.2f} (Change: +${grand_change:,.2f})")
    text_lines.append("")
    
    # Group breakdown
    text_lines.append("CONTRIBUTIONS BY ADVERTISER")
    text_lines.append("-" * 50)
    for g_id, change in sorted(group_changes.items(), key=lambda x: x[1]['current_total'], reverse=True):
        name = change['name']
        cur_tot = change['current_total']
        net_chg = change['net_change']
        text_lines.append(f"- {name}: ${cur_tot:,.2f} (Weekly Change: +${net_chg:,.2f})")
    text_lines.append("")
    
    # New Groups
    if new_groups:
        text_lines.append("NEW ADVERTISERS REGISTERED THIS WEEK")
        text_lines.append("-" * 50)
        for g in new_groups:
            text_lines.append(f"- {g['name']} (Registered: {g['registered_date']})")
        text_lines.append("")
        
    # Notable Additions
    if notable_donations:
        text_lines.append(f"NOTABLE NEW CONTRIBUTIONS (>= ${NOTABLE_THRESHOLD:,.2f})")
        text_lines.append("-" * 60)
        for d in sorted(notable_donations, key=lambda x: x['total'], reverse=True):
            text_lines.append(f"- ${d['total']:,.2f} from {d['contributor']} ({d['location']}) to {d['group_name']}")
    else:
        text_lines.append("No notable single contributions of $1,000+ this week.")
    text_lines.append("")
    
    # Summary of all new donations
    if all_new_donations:
        text_lines.append(f"TOTAL NEW CONTRIBUTIONS DETECTED: {len(all_new_donations)}")
    
    text_lines.append("")
    text_lines.append("View detailed historical CSV in the GitHub Repository.")
    body_text = "\n".join(text_lines)
    
    # ---------------------- HTML BODY (PREMIUM STYLE) ----------------------
    html_template = """<!DOCTYPE html>
<html>
<head>
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <style>
        body {{
            font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Helvetica, Arial, sans-serif;
            background-color: #f8fafc;
            color: #1e293b;
            margin: 0;
            padding: 0;
            -webkit-font-smoothing: antialiased;
        }}
        .wrapper {{
            width: 100%;
            background-color: #f8fafc;
            padding: 24px 12px;
            box-sizing: border-box;
        }}
        .container {{
            max-width: 650px;
            margin: 0 auto;
            background-color: #ffffff;
            border-radius: 12px;
            overflow: hidden;
            box-shadow: 0 4px 6px -1px rgba(0, 0, 0, 0.05), 0 2px 4px -1px rgba(0, 0, 0, 0.03);
            border: 1px solid #e2e8f0;
        }}
        .header {{
            background: linear-gradient(135deg, #1e293b 0%, #0f172a 100%);
            color: #ffffff;
            padding: 32px 24px;
            text-align: center;
        }}
        .header h1 {{
            margin: 0 0 8px 0;
            font-size: 22px;
            font-weight: 800;
            letter-spacing: -0.025em;
            color: #ffffff;
        }}
        .header .meta {{
            font-size: 14px;
            color: #94a3b8;
            margin: 0;
        }}
        .header .date-badge {{
            display: inline-block;
            background-color: #f59e0b;
            color: #0f172a;
            padding: 4px 10px;
            border-radius: 9999px;
            font-weight: 700;
            font-size: 12px;
            margin-top: 12px;
        }}
        .content {{
            padding: 24px;
        }}
        .stat-card {{
            background-color: #f1f5f9;
            border-radius: 8px;
            padding: 16px;
            margin-bottom: 24px;
            border: 1px solid #e2e8f0;
            text-align: center;
        }}
        .stat-grid {{
            display: table;
            width: 100%;
        }}
        .stat-col {{
            display: table-cell;
            width: 50%;
            padding: 0 8px;
        }}
        .stat-val {{
            font-size: 24px;
            font-weight: 800;
            color: #0f172a;
            margin-bottom: 4px;
        }}
        .stat-lbl {{
            font-size: 12px;
            color: #64748b;
            text-transform: uppercase;
            letter-spacing: 0.05em;
            margin: 0;
        }}
        .section-title {{
            font-size: 16px;
            font-weight: 700;
            color: #0f172a;
            border-bottom: 2px solid #e2e8f0;
            padding-bottom: 8px;
            margin-top: 24px;
            margin-bottom: 16px;
            text-transform: uppercase;
            letter-spacing: 0.025em;
        }}
        table {{
            width: 100%;
            border-collapse: collapse;
            margin-bottom: 24px;
        }}
        th {{
            background-color: #f8fafc;
            color: #475569;
            font-weight: 600;
            font-size: 12px;
            text-align: left;
            padding: 10px 12px;
            border-bottom: 2px solid #e2e8f0;
            text-transform: uppercase;
        }}
        td {{
            padding: 10px 12px;
            border-bottom: 1px solid #e2e8f0;
            font-size: 14px;
            color: #334155;
            vertical-align: middle;
        }}
        .text-right {{
            text-align: right;
        }}
        .positive {{
            color: #10b981;
            font-weight: 600;
        }}
        .highlight-row {{
            background-color: #fef3c7;
        }}
        .badge {{
            display: inline-block;
            padding: 2px 6px;
            border-radius: 4px;
            font-size: 11px;
            font-weight: 600;
        }}
        .badge-new {{
            background-color: #d1fae5;
            color: #065f46;
        }}
        .badge-notable {{
            background-color: #fef3c7;
            color: #92400e;
            border: 1px solid #fde68a;
        }}
        .footer {{
            background-color: #f1f5f9;
            padding: 20px 24px;
            text-align: center;
            font-size: 12px;
            color: #64748b;
            border-top: 1px solid #e2e8f0;
        }}
        .footer a {{
            color: #3b82f6;
            text-decoration: none;
        }}
    </style>
</head>
<body>
    <div class="wrapper">
        <div class="container">
            <div class="header">
                <h1>Alberta Referendum Donations Update</h1>
                <p class="meta">Weekly Campaign Contributions Report</p>
                <div class="date-badge">Data Updated: {update_date}</div>
            </div>
            
            <div class="content">
                <div class="stat-card">
                    <div class="stat-grid">
                        <div class="stat-col" style="border-right: 1px solid #cbd5e1;">
                            <div class="stat-val">${grand_total:,.2f}</div>
                            <p class="stat-lbl">Grand Total Contributions</p>
                        </div>
                        <div class="stat-col">
                            <div class="stat-val positive">+${grand_change:,.2f}</div>
                            <p class="stat-lbl">Change Since Last Run</p>
                        </div>
                    </div>
                </div>
                
                {new_groups_section}
                
                <div class="section-title">Contributions by Advertiser</div>
                <table>
                    <thead>
                        <tr>
                            <th>Advertiser Name</th>
                            <th class="text-right">Total to Date</th>
                            <th class="text-right">Change</th>
                        </tr>
                    </thead>
                    <tbody>
                        {group_rows}
                    </tbody>
                </table>
                
                <div class="section-title">Notable New Contributions</div>
                {notable_section}
                
            </div>
            
            <div class="footer">
                <p>This is an automated alert generated by the Alberta Referendum Donations Tracker.</p>
                <p>Data Source: <a href="{list_url}" target="_blank">Elections Alberta Financial Disclosure Portal</a></p>
                <p>Generated on {generation_time}</p>
            </div>
        </div>
    </div>
</body>
</html>
"""

    # Generate new groups section html
    new_groups_section = ""
    if new_groups:
        new_groups_section = """
        <div class="section-title" style="color: #059669;">New Advertisers Registered</div>
        <table>
            <thead>
                <tr>
                    <th>Advertiser Name</th>
                    <th>Registration Date</th>
                    <th>Status</th>
                </tr>
            </thead>
            <tbody>
        """
        for g in new_groups:
            new_groups_section += f"""
                <tr>
                    <td><strong>{g['name']}</strong></td>
                    <td>{g['registered_date']}</td>
                    <td><span class="badge badge-new">NEWLY REGISTERED</span></td>
                </tr>
            """
        new_groups_section += """
            </tbody>
        </table>
        """
        
    # Generate advertiser rows html
    group_rows = ""
    for g_id, change in sorted(group_changes.items(), key=lambda x: x[1]['current_total'], reverse=True):
        name = change['name']
        cur_tot = change['current_total']
        net_chg = change['net_change']
        chg_text = f"+${net_chg:,.2f}" if net_chg > 0 else "$0.00"
        chg_class = 'class="positive"' if net_chg > 0 else ""
        row_class = 'class="highlight-row"' if net_chg > 0 else ""
        
        group_rows += f"""
            <tr {row_class}>
                <td>{name}</td>
                <td class="text-right" style="font-weight: 600;">${cur_tot:,.2f}</td>
                <td class="text-right {chg_class}">{chg_text}</td>
            </tr>
        """
        
    # Generate notable contributions html
    notable_section = ""
    if notable_donations:
        notable_section = """
        <table>
            <thead>
                <tr>
                    <th>Contributor</th>
                    <th>Location</th>
                    <th>Recipient Advertiser</th>
                    <th class="text-right">Amount</th>
                </tr>
            </thead>
            <tbody>
        """
        for d in sorted(notable_donations, key=lambda x: x['total'], reverse=True):
            highlight = 'style="background-color: #fef3c7;"' if d['total'] >= 5000.0 else ""
            notable_section += f"""
                <tr {highlight}>
                    <td><strong>{d['contributor']}</strong></td>
                    <td>{d['location']}</td>
                    <td>{d['group_name']}</td>
                    <td class="text-right" style="font-weight: 700; color: #b45309;">
                        ${d['total']:,.2f}
                        { '<span class="badge badge-notable">LARGE</span>' if d['total'] >= 5000.0 else '' }
                    </td>
                </tr>
            """
        notable_section += """
            </tbody>
        </table>
        """
    else:
        notable_section = "<p style='font-size: 14px; color: #64748b; font-style: italic;'>No new single contributions of $1,000+ were detected in this update.</p>"
        
    body_html = html_template.format(
        update_date=update_date,
        grand_total=grand_total,
        grand_change=grand_change,
        new_groups_section=new_groups_section,
        group_rows=group_rows,
        notable_section=notable_section,
        list_url=LIST_URL,
        generation_time=current_time_str
    )
    
    return body_html, body_text

def main():
    parser = argparse.ArgumentParser(description="Alberta Referendum political donations scraper & tracker.")
    parser.add_argument("--force", action="store_true", help="Force updates and alerts regardless of date check.")
    parser.add_argument("--dry-run", action="store_true", help="Perform scraping and report generation but do not write files or send emails.")
    args = parser.parse_args()
    
    state = load_state()
    prev_date = state.get("last_processed_date", "")
    prev_groups = state.get("groups", {})
    
    print(f"Current local time: {datetime.now().isoformat()}")
    print(f"Loaded previous date: '{prev_date}'")
    print(f"Loaded {len(prev_groups)} previously tracked groups.")
    
    try:
        # 1. Scrape main portal list
        update_date, groups = scrape_main_page()
        print(f"Main portal update date is: '{update_date}'")
        
        # Determine if we have updates
        date_changed = update_date != prev_date
        
        new_group_ids = [gid for gid in groups if gid not in prev_groups]
        has_new_groups = len(new_group_ids) > 0
        
        should_update = date_changed or has_new_groups or args.force
        
        if not should_update:
            print("No updates detected. Exit.")
            return
            
        print("Update detected or force run enabled. Proceeding to fetch detailed XLS files...")
        
        # Load previous database records to preserve reporting dates
        old_donations = load_previous_donations()
        
        all_new_donations = []
        all_notable_donations = []
        new_groups_details = []
        group_changes = {}
        
        # New database file builder list
        new_database_rows = []
        
        # For each group, download report and parse
        for g_id, g in groups.items():
            name = g['name']
            event_id = g['event_id']
            account_id = g['account_id']
            current_total = g['total_contributions']
            
            # Identify if group is newly added
            is_new_group = g_id not in prev_groups
            if is_new_group:
                new_groups_details.append({
                    'tpa_id': g_id,
                    'name': name,
                    'registered_date': g['registered_date']
                })
                
            prev_group_data = prev_groups.get(g_id, {})
            
            if date_changed or is_new_group:
                # If a new portal date arrived, base total for this week is previous total_contributions.
                # If brand new group, base is 0.0 (or previous total_contributions if existing in prev_groups).
                prev_report_total = prev_group_data.get('total_contributions', 0.0) if not is_new_group else 0.0
            else:
                # Same portal date! Maintain previous week's base total
                if 'prev_report_total' in prev_group_data:
                    prev_report_total = prev_group_data['prev_report_total']
                elif 'total_contributions' in prev_group_data:
                    # Fallback if prev_report_total wasn't explicitly saved yet
                    prev_report_total = prev_group_data['total_contributions'] - prev_group_data.get('net_change', 0.0)
                else:
                    prev_report_total = 0.0
                    
            net_change = current_total - prev_report_total
            
            group_changes[g_id] = {
                'name': name,
                'prev_total': prev_report_total,
                'current_total': current_total,
                'net_change': net_change,
                'prev_report_total': prev_report_total
            }
            
            # Default empty lists if no account_id/event_id (rare, means $0 and not fully set up)
            parsed_donations = []
            small_total = 0.0
            
            if event_id and account_id:
                try:
                    parsed_donations, small_total = download_and_parse_xls(event_id, account_id, g_id)
                except Exception as e:
                    print(f"Error parsing XLS for group {name} ({g_id}): {e}")
                    # Fallback to empty, but keep existing ones in CSV to prevent data loss
                    # Look up from old matching database and dump back in if error occurred
                    print("Attempting to preserve existing database rows due to download error...")
                    # We will copy existing records for this group ID to new_database_rows
                    if os.path.exists(CSV_FILE):
                        with open(CSV_FILE, mode='r', encoding='utf-8') as f:
                            reader = csv.DictReader(f)
                            for row in reader:
                                if row['group_id'] == g_id:
                                    new_database_rows.append(row)
                    continue
            else:
                print(f"Group {name} ({g_id}) has no EventID or AccountID. Skipping download. (Total contributions: {current_total})")
                
            # Perform bag diff & matching to assign report dates
            matched_records, new_donations = match_donations(g_id, parsed_donations, old_donations, update_date)
            
            all_new_donations.extend(new_donations)
            
            # Identify notable single donations
            for d in new_donations:
                if d['total'] >= NOTABLE_THRESHOLD:
                    all_notable_donations.append({
                        'group_id': g_id,
                        'group_name': name,
                        'contributor': d['contributor'],
                        'location': d['location'],
                        'cash': d['cash'],
                        'valued': d['valued'],
                        'total': d['total']
                    })
                    
            # Add matched records to database rows
            for record in matched_records:
                new_database_rows.append({
                    'group_id': g_id,
                    'group_name': name,
                    'contributor': record['contributor'],
                    'location': record['location'],
                    'cash': record['cash'],
                    'valued': record['valued'],
                    'total': record['total'],
                    'report_date': record['report_date']
                })
                
            # Add small contributions summary row (if > 0)
            if small_total > 0:
                new_database_rows.append({
                    'group_id': g_id,
                    'group_name': name,
                    'contributor': "[Summary - Contributions $250 and under]",
                    'location': "",
                    'cash': 0.0,
                    'valued': 0.0,
                    'total': small_total,
                    'report_date': update_date  # Always gets current date
                })
                
        # 2. Write updated database CSV
        if not args.dry_run:
            os.makedirs(os.path.dirname(CSV_FILE), exist_ok=True)
            with open(CSV_FILE, mode='w', encoding='utf-8', newline='') as f:
                fieldnames = ['group_id', 'group_name', 'contributor', 'location', 'cash', 'valued', 'total', 'report_date']
                writer = csv.DictWriter(f, fieldnames=fieldnames)
                writer.writeheader()
                for row in new_database_rows:
                    writer.writerow(row)
            print(f"Updated unified database CSV with {len(new_database_rows)} records in {CSV_FILE}")
        else:
            print(f"[Dry Run] Would write {len(new_database_rows)} records to CSV database.")
            
        # 3. Generate and send email alert
        email_html, email_text = generate_email_report(
            update_date, groups, new_groups_details, all_notable_donations, group_changes, all_new_donations
        )
        
        subject = f"Alberta Referendum Donations Update - {update_date}"
        
        if not args.dry_run:
            send_email(subject, email_html, email_text)
        else:
            # Still output mock details in dry-run
            print("\n=== DRY RUN: Generating Mock Email ===")
            print(f"Subject: {subject}")
            print(f"Parsed {len(all_new_donations)} new contributions.")
            print(f"Parsed {len(all_notable_donations)} notable donations of $1,000+.")
            print(f"New groups added: {new_groups_details}")
            print("======================================\n")
            
        # 4. Save new state
        if not args.dry_run:
            new_groups_state = {}
            for g_id, g in groups.items():
                change_info = group_changes.get(g_id, {})
                net_chg = change_info.get('net_change', 0.0)
                prev_rep_tot = change_info.get('prev_report_total', 0.0)
                
                # If there's an update, set last_update_date to update_date
                if net_chg > 0.0 or g_id not in prev_groups:
                    last_up_date = update_date
                else:
                    last_up_date = prev_groups.get(g_id, {}).get("last_update_date", g['registered_date'])
                
                new_groups_state[g_id] = {
                    'name': g['name'],
                    'registered_date': g['registered_date'],
                    'event_id': g['event_id'],
                    'account_id': g['account_id'],
                    'total_contributions': g['total_contributions'],
                    'prev_report_total': prev_rep_tot,
                    'net_change': net_chg,
                    'last_update_date': last_up_date
                }
                
            new_state = {
                "last_processed_date": update_date,
                "groups": new_groups_state
            }
            save_state(new_state)
            
            # Export update date for GitHub Action commit message
            if "GITHUB_ENV" in os.environ:
                with open(os.environ["GITHUB_ENV"], "a") as env_file:
                    env_file.write(f"UPDATE_DATE={update_date}\n")
                    
    except Exception as e:
        print(f"Critical error during execution: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

def match_donations(group_id, parsed_donations, old_donations, current_report_date):
    """
    Match parsed new donations against old donations key queues to preserve report dates.
    """
    matched_records = []
    new_donations_list = []
    
    group_old_donations = old_donations.get(group_id, {})
    
    for item in parsed_donations:
        contributor = item['contributor']
        location = item['location']
        cash = item['cash']
        valued = item['valued']
        total = item['total']
        
        key = (contributor, location, cash, valued, total)
        
        # Match using queue pop to prevent duplicate matches
        if key in group_old_donations and len(group_old_donations[key]) > 0:
            matched_date = group_old_donations[key].pop(0)
            matched_records.append({
                'contributor': contributor,
                'location': location,
                'cash': cash,
                'valued': valued,
                'total': total,
                'report_date': matched_date
            })
        else:
            new_record = {
                'contributor': contributor,
                'location': location,
                'cash': cash,
                'valued': valued,
                'total': total,
                'report_date': current_report_date
            }
            matched_records.append(new_record)
            new_donations_list.append(new_record)
            
    return matched_records, new_donations_list

if __name__ == "__main__":
    main()
